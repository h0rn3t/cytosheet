import os
import timeit

import sys
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

try:
    from openpyxl import load_workbook as openpyxl_load_workbook
    HAS_OPENPYXL = True
except Exception:  # noqa: PIE786
    HAS_OPENPYXL = False

from cytosheet import Workbook, load_workbook


def test_parse_xlsx():
    test_file = os.path.join(os.path.dirname(__file__), "test.xlsx")
    wb = load_workbook(test_file)
    print(wb.sheetnames)
    # assert "Sheet1" in ['Sheet', 'sheet1']
    assert  ['Sheet', 'sheet1'] ==  wb.sheetnames


def test_read_cells():
    test_file = os.path.join(os.path.dirname(__file__), "file_example_XLSX_5000.xlsx")

    start = timeit.default_timer()
    wb = load_workbook(test_file)
    ws = wb["Sheet1"]
    assert ws["A2"].value == "1"
    cytosheet_time = timeit.default_timer() - start
    print(f"cytosheet Time: {cytosheet_time:.4f} seconds")
    if HAS_OPENPYXL:
        start = timeit.default_timer()
        wb = openpyxl_load_workbook(test_file)
        ws = wb["Sheet1"]
        assert str(ws["A2"].value) == "1"
        openpyxl_time = timeit.default_timer() - start
        print(f"openpyxl Time: {openpyxl_time:.4f} seconds")


def test_write_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"].value = "Hello"
    ws["B2"].value = "World"

    assert ws["A1"].value == "Hello"
    assert ws["B2"].value == "World"

    file_path = tmp_path / "cytosheet_create_file_test.xlsx"
    wb.save(str(file_path))

    wb2 = load_workbook(str(file_path))
    ws2 = wb2.active
    assert ws2["A1"].value == "Hello"
    assert ws2["B2"].value == "World"


def test_merge_and_formula(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A1:B2")
    ws["A1"].value = "Test"
    ws["A1"].formula = "=SUM(1,1)"

    file_path = tmp_path / "merge_formula.xlsx"
    wb.save(str(file_path))

    wb2 = load_workbook(str(file_path))
    ws2 = wb2.active
    assert "A1:B2" in ws2.merged_cells
    assert ws2["A1"].value == "Test"
    assert ws2["A1"].formula == "=SUM(1,1)"


def test_append_and_max():
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B", "C"])
    ws.append([1, 2, 3])
    assert ws.max_row == 2
    assert ws.max_column == 3
    assert ws["A2"].value == 1
    assert ws["C1"].value == "C"


def test_iter_rows_and_cols():
    wb = Workbook()
    ws = wb.active
    ws["A1"].value = "A1"
    ws["B1"].value = "B1"
    ws["A2"].value = "A2"
    ws["B2"].value = "B2"

    rows = list(ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=2))
    assert rows[0][0].value == "A1"
    assert rows[1][1].value == "B2"

    cols = list(ws.iter_cols(min_col=1, max_col=2, min_row=1, max_row=2))
    assert cols[0][1].value == "A2"
    assert cols[1][0].value == "B1"
