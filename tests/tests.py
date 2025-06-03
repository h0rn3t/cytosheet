import os
import timeit
from openpyxl import load_workbook as openpyxl_load_workbook, Workbook as OpenpyxlWorkbook
import pytest
from src.cytosheet import Workbook, load_workbook


import sys
sys.path.append("src")

def test_parse_xlsx():
    """Тестирование парсинга существующего XLSX файла."""
    test_file = os.path.join(os.path.dirname(__file__), 'test.xlsx')

    wb = load_workbook(test_file)

    assert 'sheet1' == wb.get_sheet_by_name('Sheet1').title


def test_read_cells():
    """Тестирование чтения значений ячеек."""
    test_file = os.path.join(os.path.dirname(__file__), 'file_example_XLSX_5000.xlsx')

    # Тестирование cytosheet
    start = timeit.default_timer()
    wb = load_workbook(test_file)
    ws = wb.get_sheet_by_name('Sheet1')
    assert ws['A2'].value == 1  # Приведение типов, если нужно
    stop = timeit.default_timer()
    cytosheet_time = stop - start
    print(f'\ncytosheet Time: {cytosheet_time:.4f} seconds')

    # Тестирование openpyxl
    start = timeit.default_timer()
    wb = openpyxl_load_workbook(test_file)
    ws = wb.get_sheet_by_name('Sheet1')
    assert str(ws['A2'].value) == '1'  # Приведение типов
    stop = timeit.default_timer()
    openpyxl_time = stop - start
    print(f'\nopenpyxl Time: {openpyxl_time:.4f} seconds')


def test_write_cells():
    """Тестирование записи значений ячеек."""

    # Создаем эксельку, и записываем данные, сохраняем
    wb = Workbook() # по умолчанию создается один лист с названием 'Sheet'
    # wb = OpenpyxlWorkbook()
    ws = wb.active

    ws['A1'].value = 'Hello'
    ws['B2'].value = 'World'

    assert ws['A1'].value == 'Hello'
    assert ws['B2'].value == 'World'

    file_path = os.path.join(os.path.dirname(__file__), 'cytosheet_create_file_test.xlsx')

    wb.save(file_path)

    # Открываем созданный файл и проверяем значения
    # wb2 = OpenpyxlWorkbook()
    wb2 = load_workbook(file_path)
    ws2 = wb2.active

    assert ws2['A1'].value == 'Hello'
    assert ws2['B2'].value == 'World'

    # Удаляем файл после теста
    os.remove(file_path)


def test_numeric_cells():
    wb = Workbook()
    ws = wb.active

    ws['A1'] = 42
    ws['B1'] = 3.14
    ws['C1'] = 'text'

    file_path = os.path.join(os.path.dirname(__file__), 'numeric_test.xlsx')
    wb.save(file_path)

    wb2 = load_workbook(file_path)
    ws2 = wb2.active

    assert ws2['A1'].value == 42
    assert ws2['B1'].value == 3.14
    assert ws2['C1'].value == 'text'

    os.remove(file_path)


# Потоковое чтение строк (lazy=True) без загрузки листа целиком.
def test_lazy_iter_rows():
    """Потоковое чтение строк (lazy=True) без загрузки листа целиком."""
    test_file = os.path.join(os.path.dirname(__file__),
                             'file_example_XLSX_5000.xlsx')

    wb = load_workbook(test_file, lazy=True)   # новый режим
    ws = wb.get_sheet_by_name('Sheet1')

    rows = ws.iter_rows(values_only=True)

    header = next(rows)            # первая строка – заголовок
    data   = next(rows)            # вторая – реальные данные

    # В файле примерного набора значение 1 находится в A2
    assert data[0] == 1

    wb.close()
