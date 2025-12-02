import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))

from openpyxl import Workbook as OpenpyxlWorkbook
from cytosheet import Workbook as CytosheetWorkbook, Font, Alignment, Border, Side, PatternFill, Protection, Color


def _create_style_objects():
    font = Font(name='Arial', size=12, bold=True)
    alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    border = Border(
        left=Side(style='thin', color=Color(rgb='FF0000')),
        right=Side(style='thin', color=Color(rgb='00FF00')),
        top=Side(style='thin', color=Color(rgb='0000FF')),
        bottom=Side(style='thin', color=Color(rgb='FFFF00')),
    )
    fill = PatternFill(patternType='solid', fgColor=Color(rgb='FFFF00'))
    protection = Protection(locked=True, hidden=False)
    return font, alignment, border, fill, protection


def test_cell_font_proxy_sets_and_gets_style_instance():
    wb = CytosheetWorkbook()
    ws = wb.active
    cell = ws['A1']

    font, _, _, _, _ = _create_style_objects()

    # В cytosheet, как и в openpyxl, у новой ячейки уже есть дефолтный Style
    assert cell.style is not None
    cell.font = font

    assert cell.style is not None
    assert cell.font is font
    assert cell.style.font is font


def test_cell_alignment_proxy_sets_and_gets_style_instance():
    wb = CytosheetWorkbook()
    ws = wb.active
    cell = ws['A1']

    _, alignment, _, _, _ = _create_style_objects()

    assert cell.style is not None
    cell.alignment = alignment

    assert cell.style is not None
    assert cell.alignment is alignment
    assert cell.style.alignment is alignment


def test_cell_border_proxy_sets_and_gets_style_instance():
    wb = CytosheetWorkbook()
    ws = wb.active
    cell = ws['A1']

    _, _, border, _, _ = _create_style_objects()

    assert cell.style is not None
    cell.border = border

    assert cell.style is not None
    assert cell.border is border
    assert cell.style.border is border


def test_cell_fill_proxy_sets_and_gets_style_instance():
    wb = CytosheetWorkbook()
    ws = wb.active
    cell = ws['A1']

    _, _, _, fill, _ = _create_style_objects()

    assert cell.style is not None
    cell.fill = fill

    assert cell.style is not None
    assert cell.fill is fill
    assert cell.style.fill is fill


def test_cell_protection_proxy_sets_and_gets_style_instance():
    wb = CytosheetWorkbook()
    ws = wb.active
    cell = ws['A1']

    _, _, _, _, protection = _create_style_objects()

    assert cell.style is not None
    cell.protection = protection

    assert cell.style is not None
    assert cell.protection is protection
    assert cell.style.protection is protection


def test_style_proxies_roundtrip_through_file(tmp_path):
    """Записываем стили через прокси, читаем openpyxl и cytosheet и сравниваем базовые поля."""
    wb = CytosheetWorkbook()
    ws = wb.active
    font, alignment, border, fill, protection = _create_style_objects()

    c = ws['A1']
    c.value = 'styled'
    c.font = font
    c.alignment = alignment
    c.border = border
    c.fill = fill
    c.protection = protection

    file_path = tmp_path / 'style_proxies_roundtrip.xlsx'
    wb.save(str(file_path))

    # openpyxl читает файл без ошибок
    from openpyxl import load_workbook as oxl_load_workbook

    oxl_wb = oxl_load_workbook(str(file_path))
    oxl_ws = oxl_wb.active
    oxl_cell = oxl_ws['A1']

    assert oxl_cell.value == 'styled'

    # Минимальные проверки того, что стили применились
    assert oxl_cell.font is not None
    assert oxl_cell.alignment is not None
    assert oxl_cell.border is not None
    assert oxl_cell.fill is not None
    assert oxl_cell.protection is not None

    # Cytosheet тоже корректно читает свои же стили
    from cytosheet import load_workbook as cyto_load_workbook

    wb2 = cyto_load_workbook(str(file_path))
    ws2 = wb2.active
    c2 = ws2['A1']

    assert c2.value == 'styled'
    assert c2.font is not None
    assert c2.alignment is not None
    assert c2.border is not None
    assert c2.fill is not None
    assert c2.protection is not None
