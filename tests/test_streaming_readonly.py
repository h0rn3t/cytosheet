"""Тести потокового read-only читання та lazy-збереження (Фаза 3, task 5.1, D-10).

Кожен тест = один `#### Scenario:` зі specs зміни full-library-roadmap:
- `xlsx-io`: lazy не читає XML листа при відкритті; save на read-only книзі;
  save після close → ValueError; емісія `<dimension>`
- `worksheet-operations`: позиції за `r=`, падінг розріджених рядків, фолбек без
  `<dimension>`, identity комірок у звичайному режимі
"""

import os
import sys
import zipfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

import pytest
from openpyxl import load_workbook as oxl_load

from cytosheet import Workbook, load_workbook


class SpyZipFile(zipfile.ZipFile):
    """ZipFile, що рахує звернення до `read` — для перевірки, що lazy не читає лист."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.read_calls = []

    def read(self, name, *args, **kwargs):
        self.read_calls.append(name if isinstance(name, str) else name.filename)
        return super().read(name, *args, **kwargs)


def _make_file(path, sheets):
    """Створити .xlsx через openpyxl: sheets = {назва: [[рядок], ...]}."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return str(path)


@pytest.fixture
def simple_path(tmp_path):
    return _make_file(tmp_path / "simple.xlsx", {
        "S1": [["a", 1], ["b", 2], ["c", 3]],
    })


# --- xlsx-io: Lazy-режим не читає XML листа при відкритті -------------------

def test_lazy_open_does_not_read_sheet_xml(simple_path):
    """Scenario: Відкриття в lazy-режимі не звертається до XML листа."""
    archive = SpyZipFile(simple_path, 'r')
    Workbook(_archive=archive, lazy=True)

    sheet_reads = [c for c in archive.read_calls if c.startswith('xl/worksheets/')]
    assert sheet_reads == [], f"lazy прочитав XML листа: {sheet_reads}"


def test_eager_open_reads_sheet_xml_exactly_once(simple_path):
    """Scenario: Звичайне відкриття читає XML листа рівно один раз."""
    archive = SpyZipFile(simple_path, 'r')
    Workbook(_archive=archive, lazy=False)

    sheet_reads = [c for c in archive.read_calls if c.startswith('xl/worksheets/')]
    assert len(sheet_reads) == len(set(sheet_reads)), f"подвійне читання: {sheet_reads}"
    assert sheet_reads.count('xl/worksheets/sheet1.xml') == 1


# --- xlsx-io: Збереження книги, відкритої в read-only/lazy режимі -----------

def test_lazy_save_preserves_all_sheets(tmp_path):
    """Scenario: lazy без змін → save зберігає дані всіх листів."""
    src = _make_file(tmp_path / "multi.xlsx", {
        "One": [["one-a1", 11]],
        "Two": [["two-a1", 22]],
        "Three": [["three-a1", 33]],
    })
    out = str(tmp_path / "out.xlsx")

    wb = load_workbook(src, read_only=True)
    wb.save(out)

    oxl = oxl_load(out)
    assert oxl["One"]["A1"].value == "one-a1"
    assert oxl["One"]["B1"].value == 11
    assert oxl["Two"]["A1"].value == "two-a1"
    assert oxl["Three"]["A1"].value == "three-a1"
    assert oxl["Three"]["B1"].value == 33


def test_lazy_modify_save_applies_change_and_keeps_rest(tmp_path):
    """Scenario: lazy зі зміною → save застосовує зміну й зберігає решту."""
    src = _make_file(tmp_path / "multi.xlsx", {
        "One": [["one-a1", 11]],
        "Two": [["two-a1", 22]],
    })
    out = str(tmp_path / "out.xlsx")

    wb = load_workbook(src, read_only=True)
    wb["One"]["A1"] = "CHANGED"
    wb.save(out)

    oxl = oxl_load(out)
    assert oxl["One"]["A1"].value == "CHANGED"
    assert oxl["Two"]["A1"].value == "two-a1"
    assert oxl["Two"]["B1"].value == 22


# --- xlsx-io: Збереження після закриття архіву ------------------------------

def test_save_after_close_raises(simple_path, tmp_path):
    """Scenario: save після close підіймає ValueError."""
    wb = load_workbook(simple_path)
    wb.close()

    with pytest.raises(ValueError):
        wb.save(str(tmp_path / "out.xlsx"))


def test_save_new_workbook_without_archive_still_works(tmp_path):
    """Нова книга не має архіву — це не «закрита» книга, save має працювати."""
    wb = Workbook()
    wb.active["A1"] = "hello"
    out = str(tmp_path / "new.xlsx")
    wb.save(out)

    assert oxl_load(out).active["A1"].value == "hello"


# --- xlsx-io: Емісія <dimension> --------------------------------------------

def test_dimension_emitted_on_write(tmp_path):
    """Scenario: Записаний cytosheet файл стримиться без фолбеку."""
    out = str(tmp_path / "dim.xlsx")
    wb = Workbook()
    ws = wb.active
    for row in [["a", "b", "c"], [1, 2, 3], ["x", "y", "z"]]:
        ws.append(row)
    wb.save(out)

    xml = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode()
    assert '<dimension ref="A1:C3"/>' in xml

    wb2 = load_workbook(out, read_only=True)
    rows = list(wb2.active.iter_rows(values_only=True))
    assert all(len(r) == 3 for r in rows), rows
    assert wb2.active._preloaded is False, "лист матеріалізувався замість стримінгу"


# --- worksheet-operations: позиції за r= -----------------------------------

def _sparse_path(tmp_path):
    """Лист, де B1 відсутня у XML: заповнені лише A1 і C1."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 10
    ws["C1"] = 30
    path = str(tmp_path / "sparse.xlsx")
    wb.save(path)
    return path


def test_stream_sparse_row_keeps_column_positions(tmp_path):
    """Scenario: Розріджений рядок зберігає позиції колонок."""
    path = _sparse_path(tmp_path)

    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    assert 'r="B1"' not in xml, "фікстура не розріджена — тест нічого не доводить"

    ws = load_workbook(path, read_only=True).active
    row = next(iter(ws.iter_rows(values_only=True)))
    assert row == (10, None, 30)


def test_stream_sparse_matches_openpyxl(tmp_path):
    """Scenario (AND): той самий файл через openpyxl read_only дає той самий кортеж."""
    path = _sparse_path(tmp_path)

    cyto_row = next(iter(load_workbook(path, read_only=True).active.iter_rows(values_only=True)))
    oxl_row = next(iter(oxl_load(path, read_only=True).active.iter_rows(values_only=True)))
    assert cyto_row == oxl_row


def test_stream_missing_row_yields_nones(tmp_path):
    """Scenario: Пропущений рядок віддається кортежем з None."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "r1"
    ws["A3"] = "r3"
    path = str(tmp_path / "gap.xlsx")
    wb.save(path)

    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    assert 'r="2"' not in xml, "фікстура має рядок 2 — тест нічого не доводить"

    rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[1] == (None,)
    assert rows[0] == ("r1",) and rows[2] == ("r3",)


# --- worksheet-operations: фолбек без <dimension> ---------------------------

def test_stream_no_dimension_falls_back(tmp_path):
    """Scenario: Лист без dimension читається коректно (через матеріалізацію)."""
    import openpyxl
    import re

    src = _make_file(tmp_path / "src.xlsx", {"S1": [["a", 1], ["b", 2]]})
    stripped = str(tmp_path / "nodim.xlsx")

    # Перебираємо архів, вирізаючи <dimension .../> з XML листа
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(stripped, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            zout.writestr(item, data)

    assert b"<dimension" not in zipfile.ZipFile(stripped).read("xl/worksheets/sheet1.xml")

    ws = load_workbook(stripped, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))

    assert ws._preloaded is True, "без <dimension> мав спрацювати фолбек на матеріалізацію"
    widths = {len(r) for r in rows}
    assert widths == {ws.max_column}, f"рвані кортежі: {rows}"
    assert rows == list(oxl_load(stripped).active.iter_rows(values_only=True))


# --- worksheet-operations: identity комірок у звичайному режимі --------------

def test_iter_rows_identity_in_eager_mode(tmp_path, simple_path):
    """Scenario: Запис через ітератор видно в листі."""
    out = str(tmp_path / "eager.xlsx")

    wb = load_workbook(simple_path)  # без read_only
    ws = wb.active
    for row in ws.iter_rows():
        row[0].value = "X"
        break

    assert ws["A1"].value == "X"

    wb.save(out)
    assert oxl_load(out).active["A1"].value == "X"


def test_stream_and_materialized_paths_agree(tmp_path):
    """Обидві гілки ітерації дають однакові кортежі на одному файлі (design risk)."""
    path = _make_file(tmp_path / "both.xlsx", {
        "S1": [["a", 1, None], [None, 2, "z"], ["c", None, 3]],
    })

    streamed = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    materialized = list(load_workbook(path).active.iter_rows(values_only=True))

    assert streamed == materialized
    assert streamed == list(oxl_load(path).active.iter_rows(values_only=True))


# --- worksheet-operations: пікова памʼять не залежить від кількості рядків ---

def _write_big_xlsx(path, n_rows):
    """Мінімальний валідний .xlsx на n_rows рядків, зібраний сирим XML.

    Навмисно не через openpyxl: той будує всі рядки в памʼяті й на 200k рядків
    сам стає вузьким місцем тесту. Тут — 0.1 с і повний контроль над `<dimension>`,
    від якого залежить, чи ввімкнеться стримінг.
    """
    rows = "".join(
        f'<row r="{r}"><c r="A{r}" t="inlineStr"><is><t>рядок-{r}</t></is></c>'
        f'<c r="B{r}"><v>{r}</v></c></row>'
        for r in range(1, n_rows + 1)
    )
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:B{n_rows}"/><sheetData>{rows}</sheetData></worksheet>'
    )
    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                   '</Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                   '</Relationships>')
        z.writestr("xl/workbook.xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                   'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<sheets><sheet name="Big" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                   '</Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", sheet)
    return str(path)


_PEAK_RSS_SCRIPT = """
import resource, sys
sys.path.insert(0, %r)
from cytosheet import load_workbook

ws = load_workbook(sys.argv[1], read_only=True).active
count = 0
last = None
for row in ws.iter_rows(values_only=True):
    count += 1
    last = row
assert not ws._preloaded, "лист матеріалізувався — вимір памʼяті не про стримінг"
print(count, last[0], resource.getrusage(resource.RUSAGE_SELF).ru_maxrss)
"""


def _stream_peak_rss(path):
    """Пікова RSS окремого процесу, що стримить файл цілком.

    Окремий процес принципово: ru_maxrss монотонний, тож два заміри в одному
    процесі дали б пік першого як підлогу для другого.
    """
    import subprocess

    src = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src'))
    out = subprocess.run(
        [sys.executable, "-c", _PEAK_RSS_SCRIPT % src, path],
        capture_output=True, text=True, check=True,
    )
    count, last_value, rss = out.stdout.split()
    return int(count), last_value, int(rss)


def test_stream_peak_memory_constant_in_rows(tmp_path):
    """Scenario: Пікова памʼять не залежить від кількості рядків."""
    small = _write_big_xlsx(tmp_path / "rows_10k.xlsx", 10_000)
    big = _write_big_xlsx(tmp_path / "rows_200k.xlsx", 200_000)

    small_count, small_last, small_rss = _stream_peak_rss(small)
    big_count, big_last, big_rss = _stream_peak_rss(big)

    # Дані справді прочитані до кінця, інакше "стала памʼять" нічого не варта
    assert (small_count, small_last) == (10_000, "рядок-10000")
    assert (big_count, big_last) == (200_000, "рядок-200000")

    growth = (big_rss - small_rss) / small_rss
    assert growth < 0.20, (
        f"пік RSS виріс на {growth:.0%} за 20x даних "
        f"({small_rss} → {big_rss}); стримінг тримає рядки в памʼяті"
    )
