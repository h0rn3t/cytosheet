"""Вірність типів даних при round-trip (Фаза 2 зміни full-library-roadmap).

Кожен тест = один `#### Scenario:` зі спеки `cell-model`:
- D-7: bool ↔ bool (а не 1/0)
- D-8: дати/час ↔ datetime за датовим number_format (система 1900)
- D-9: inlineStr читається кожною стратегією парсингу
- 4.4: значення-помилки t="e" зберігають data_type='e'

Ground truth — openpyxl: перевіряємо не «не None», а точний збіг значень і типів.
"""

import datetime
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

import pytest
from openpyxl import Workbook as OxWorkbook, load_workbook as oxl_load

from cytosheet import Workbook, load_workbook


def test_bool_roundtrip_keeps_type(tmp_path):
    """Scenario: Boolean round-trip зберігає тип."""
    out = str(tmp_path / "bool.xlsx")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = True
    ws["A2"] = False
    ws["A3"] = 1  # int поруч із bool не має злитися з True
    wb.save(out)

    ws2 = load_workbook(out).active
    assert ws2["A1"].value is True
    assert ws2["A2"].value is False
    assert isinstance(ws2["A3"].value, int) and not isinstance(ws2["A3"].value, bool)

    # openpyxl бачить те саме
    o = oxl_load(out).active
    assert (o["A1"].value, o["A2"].value, o["A3"].value) == (True, False, 1)


@pytest.mark.parametrize("coord,value", [
    ("A1", datetime.datetime(2025, 1, 2)),
    ("A2", datetime.datetime(2025, 1, 2, 15, 30, 45)),
    ("A3", datetime.datetime(1900, 1, 5)),   # до міфічного 1900-02-29 (serial < 60)
    ("A4", datetime.datetime(1900, 3, 1)),   # після нього
    ("A5", datetime.time(9, 15)),
])
def test_datetime_roundtrip_matches_openpyxl(tmp_path, coord, value):
    """Scenario: Дата round-trip повертає datetime."""
    ours = str(tmp_path / "d_cyto.xlsx")
    theirs = str(tmp_path / "d_oxl.xlsx")

    wb = Workbook()
    wb.active[coord] = value
    wb.save(ours)

    # Еталон: те саме значення, записане openpyxl
    ox = OxWorkbook()
    ox.active[coord] = value
    ox.save(theirs)

    expected = oxl_load(theirs).active[coord].value
    assert load_workbook(ours).active[coord].value == expected   # cyto→cyto
    assert oxl_load(ours).active[coord].value == expected        # cyto→openpyxl


def test_date_read_from_openpyxl_file(tmp_path):
    """Scenario: Дата round-trip повертає datetime (файл написаний openpyxl)."""
    path = str(tmp_path / "oxl_dates.xlsx")
    ox = OxWorkbook()
    ox.active["A1"] = datetime.datetime(2025, 6, 17, 12, 0)
    ox.active["A2"] = datetime.time(9, 15)     # builtin numFmtId=21, не custom numFmt
    ox.active["A3"] = 42                       # число з тим самим типом <v> лишається числом
    ox.save(path)

    ws = load_workbook(path).active
    assert ws["A1"].value == datetime.datetime(2025, 6, 17, 12, 0)
    assert ws["A2"].value == datetime.time(9, 15)
    assert ws["A3"].value == 42

    # Потоковий шлях не має розходитись із матеріалізуючим
    lazy = load_workbook(path, read_only=True).active
    assert list(lazy.iter_rows(values_only=True)) == \
           list(oxl_load(path, read_only=True).active.iter_rows(values_only=True))


def test_error_value_keeps_data_type(tmp_path):
    """Scenario: значення-помилка t="e" зберігає тип помилки."""
    path = str(tmp_path / "err.xlsx")
    ox = OxWorkbook()
    ox.active["A1"] = "#DIV/0!"
    ox.save(path)

    cell = load_workbook(path).active["A1"]
    ref = oxl_load(path).active["A1"]
    assert cell.value == ref.value == "#DIV/0!"
    assert cell.data_type == ref.data_type == "e"


@pytest.mark.parametrize("text", [
    "текст",              # не-ASCII: openpyxl пише як &#1090;&#1077;...
    "a & b < c",          # &amp; / &lt;
    'x > y & "q"',
])
def test_xml_entities_decoded_by_simple_parser(tmp_path, text):
    """Текст із XML-сутностями не має протікати сирим у значення комірки.

    Малі файли йдуть простим string-парсером, який ріже сирий підрядок; без
    зняття сутностей 'a & b' повертався б як 'a &amp; b' (тиха порча даних).
    """
    path = str(tmp_path / "ents.xlsx")
    ox = OxWorkbook()
    ox.active["A1"] = text
    ox.save(path)

    ws = load_workbook(path).active
    assert ws._is_small_file, "файл великий — тест не перевіряє простий парсер"
    assert ws["A1"].value == text == oxl_load(path).active["A1"].value
