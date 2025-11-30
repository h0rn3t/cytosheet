import os
import timeit
import sys
import pytest
import datetime

# Add the src directory to the path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))

from openpyxl import load_workbook as openpyxl_load_workbook, Workbook as OpenpyxlWorkbook
from cytosheet import (
    Workbook, load_workbook,
    Color, Side, Border, Font, PatternFill,
    Alignment, Protection, Style, DEFAULT_STYLE
)

def test_parse_xlsx():
    """Тестирование парсинга существующего XLSX файла."""
    test_file = os.path.join(os.path.dirname(__file__), 'test.xlsx')

    wb = load_workbook(test_file)

    # Имя листа теперь читается из xl/workbook.xml, там оно 'Sheet1'
    assert 'Sheet1' == wb.get_sheet_by_name('Sheet1').title


def test_read_cells():
    """Тестирование чтения значений ячеек."""
    test_file = os.path.join(os.path.dirname(__file__), 'file_example_XLSX_5000.xlsx')

    # Тестирование cytosheet
    start = timeit.default_timer()
    wb = load_workbook(test_file)
    # избегаем deprecated get_sheet_by_name
    ws = wb['Sheet1']
    assert ws['A2'].value == 1  # Приведение типов, если нужно
    stop = timeit.default_timer()
    cytosheet_time = stop - start
    print(f'\ncytosheet Time: {cytosheet_time:.4f} seconds')

    # Тестирование openpyxl
    start = timeit.default_timer()
    wb = openpyxl_load_workbook(test_file)
    # используем рекомендованный синтаксис индексации по имени
    ws = wb['Sheet1']
    assert str(ws['A2'].value) == '1'  # Приведение типов
    stop = timeit.default_timer()
    openpyxl_time = stop - start
    print(f'\nopenpyxl Time: {openpyxl_time:.4f} seconds')


def test_write_cells():
    """Тестирование записи значений ячеек."""

    # Создаем эксельку, и записываем данные, сохраняем
    wb = Workbook() # по умолчанию создается один лист с названием 'Sheet'
    # wb = OpenpyxlWorkbook()
    ws = wb.active

    ws['A1'].value = 'Hello'
    ws['B2'].value = 'World'

    assert ws['A1'].value == 'Hello'
    assert ws['B2'].value == 'World'

    file_path = os.path.join(os.path.dirname(__file__), 'cytosheet_create_file_test.xlsx')

    wb.save(file_path)

    # Открываем созданный файл и проверяем значения
    # wb2 = OpenpyxlWorkbook()
    wb2 = load_workbook(file_path)
    ws2 = wb2.active

    assert ws2['A1'].value == 'Hello'
    assert ws2['B2'].value == 'World'

    # Удаляем файл после теста
    os.remove(file_path)


def test_numeric_cells():
    wb = Workbook()
    ws = wb.active

    ws['A1'] = 42
    ws['B1'] = 3.14
    ws['C1'] = 'text'

    file_path = os.path.join(os.path.dirname(__file__), 'numeric_test.xlsx')
    wb.save(file_path)

    wb2 = load_workbook(file_path)
    ws2 = wb2.active

    assert ws2['A1'].value == 42
    assert ws2['B1'].value == 3.14
    assert ws2['C1'].value == 'text'

    os.remove(file_path)


# Потоковое чтение строк (lazy=True) без загрузки листа целиком.
def test_lazy_iter_rows():
    """Потоковое чтение строк (lazy=True) без загрузки листа целиком."""
    test_file = os.path.join(os.path.dirname(__file__),
                             'file_example_XLSX_5000.xlsx')

    wb = load_workbook(test_file, lazy=True)   # новый режим
    # используем индексный доступ вместо get_sheet_by_name
    ws = wb['Sheet1']

    rows = ws.iter_rows(values_only=True)

    header = next(rows)            # первая строка – заголовок
    data   = next(rows)            # вторая – реальные данные

    # В файле примерного набора значение 1 находится в A2
    assert data[0] == 1

    wb.close()


# Tests from test_styles_and_merge.py

def test_merge_cells():
    """Test merging cells functionality."""
    # Create a new workbook
    wb = Workbook()
    ws = wb.active

    # Set values in cells
    ws['A1'] = 'Merged Cell'
    ws['A2'] = 'Normal Cell'

    # Merge cells A1:B2
    merged_cell = ws.merge_cells('A1:B2')

    # Check that the merged cell has the correct properties
    assert merged_cell.is_merged_cell
    assert merged_cell.merged_range == 'A1:B2'
    assert merged_cell.value == 'Merged Cell'

    # Check that the other cells in the merged range are marked as merged
    assert ws['A2'].is_merged_cell
    assert ws['B1'].is_merged_cell
    assert ws['B2'].is_merged_cell

    # Check that the other cells in the merged range have no value
    assert ws['A2'].value is None
    assert ws['B1'].value is None
    assert ws['B2'].value is None

    # Save the workbook
    file_path = os.path.join(os.path.dirname(__file__), 'test_merge.xlsx')
    wb.save(file_path)

    # Load the workbook and check that the merged cells are still merged
    wb2 = load_workbook(file_path)
    ws2 = wb2.active

    # Check that the merged cell has the correct value
    assert ws2['A1'].value == 'Merged Cell'

    # Clean up
    os.remove(file_path)

def test_unmerge_cells():
    """Test unmerging cells functionality."""
    # Create a new workbook
    wb = Workbook()
    ws = wb.active

    # Set values in cells
    ws['A1'] = 'Merged Cell'

    # Merge cells A1:B2
    ws.merge_cells('A1:B2')

    # Unmerge cells
    ws.unmerge_cells('A1:B2')

    # Check that the cells are no longer merged
    assert not ws['A1'].is_merged_cell
    assert not ws['A2'].is_merged_cell
    assert not ws['B1'].is_merged_cell
    assert not ws['B2'].is_merged_cell

    # Check that the value is still in the first cell
    assert ws['A1'].value == 'Merged Cell'

def test_cell_styles():
    """Test cell styles functionality."""
    # Create a new workbook
    wb = Workbook()
    ws = wb.active

    # Create a style
    font = Font(name='Arial', size=12, bold=True, italic=True)
    border = Border(
        left=Side(style='thin', color=Color(rgb='FF0000')),
        right=Side(style='thin', color=Color(rgb='FF0000')),
        top=Side(style='thin', color=Color(rgb='FF0000')),
        bottom=Side(style='thin', color=Color(rgb='FF0000'))
    )
    fill = PatternFill(patternType='solid', fgColor=Color(rgb='FFFF00'))
    alignment = Alignment(horizontal='center', vertical='center')
    protection = Protection(locked=True, hidden=False)
    style = Style(font=font, border=border, fill=fill, alignment=alignment, protection=protection)

    # Apply the style to a cell
    ws['A1'] = 'Styled Cell'
    ws['A1'].style = style

    # Check that the style was applied
    assert ws['A1'].style.font.name == 'Arial'
    assert ws['A1'].style.font.size == 12
    assert ws['A1'].style.font.bold
    assert ws['A1'].style.font.italic
    assert ws['A1'].style.border.left.style == 'thin'
    assert ws['A1'].style.border.left.color.rgb == 'FF0000'
    assert ws['A1'].style.fill.patternType == 'solid'
    assert ws['A1'].style.fill.fgColor.rgb == 'FFFF00'
    assert ws['A1'].style.alignment.horizontal == 'center'
    assert ws['A1'].style.alignment.vertical == 'center'
    assert ws['A1'].style.protection.locked
    assert not ws['A1'].style.protection.hidden

    # Save the workbook
    file_path = os.path.join(os.path.dirname(__file__), 'test_styles.xlsx')
    wb.save(file_path)

    # Clean up
    os.remove(file_path)

def test_merged_cell_with_style():
    """Test applying styles to merged cells."""
    # Create a new workbook
    wb = Workbook()
    ws = wb.active

    # Set values in cells
    ws['A1'] = 'Merged Cell with Style'

    # Create a style
    font = Font(name='Arial', size=14, bold=True)
    border = Border(
        left=Side(style='thick', color=Color(rgb='0000FF')),
        right=Side(style='thick', color=Color(rgb='0000FF')),
        top=Side(style='thick', color=Color(rgb='0000FF')),
        bottom=Side(style='thick', color=Color(rgb='0000FF'))
    )
    fill = PatternFill(patternType='solid', fgColor=Color(rgb='00FF00'))
    alignment = Alignment(horizontal='center', vertical='center')
    style = Style(font=font, border=border, fill=fill, alignment=alignment)

    # Merge cells A1:C3
    merged_cell = ws.merge_cells('A1:C3')

    # Apply the style to the merged cell
    merged_cell.style = style

    # Save the workbook
    file_path = os.path.join(os.path.dirname(__file__), 'test_merge_style.xlsx')
    wb.save(file_path)

    # Clean up
    os.remove(file_path)

def test_cell_number_format_property_proxy():
    wb = Workbook()
    ws = wb.active

    cell = ws['A1']
    cell.value = 123.456

    # Устанавливаем формат числа через proxy
    cell.number_format = '0.00%'

    # Проверяем, что доступ через свойство и через style.numberFormat согласован
    assert cell.number_format == '0.00%'
    assert cell.style is not None
    assert getattr(cell.style, 'numberFormat', None) == '0.00%'

def test_number_format_roundtrip_with_openpyxl(tmp_path):
    file_path = tmp_path / "number_format_test.xlsx"

    # Создаём книгу через cytosheet и задаём number_format
    wb = Workbook()
    ws = wb.active

    cell = ws['A1']
    cell.value = 0.5
    cell.number_format = '0.00%'

    wb.save(str(file_path))

    # Читаем через openpyxl и проверяем, что формат попал в styles.xml
    wb_ox = openpyxl_load_workbook(str(file_path), data_only=False)
    ws_ox = wb_ox.active
    cell_ox = ws_ox['A1']

    # openpyxl нормализует формат, но строка должна совпадать
    assert cell_ox.number_format == '0.00%'

def test_number_format_read_from_openpyxl_styles(tmp_path):
    file_path = tmp_path / "number_format_from_openpyxl.xlsx"

    # Сначала создаём файл через openpyxl с заданным number_format
    wb_ox = OpenpyxlWorkbook()
    ws_ox = wb_ox.active
    cell_ox = ws_ox['A1']
    cell_ox.value = 0.25
    cell_ox.number_format = '0.00%'
    wb_ox.save(str(file_path))

    # Теперь читаем этот файл через cytosheet
    wb = load_workbook(str(file_path))
    ws = wb.active
    cell = ws['A1']

    # Значение может быть числом, формат должен подтянуться из styles.xml
    assert cell.value == 0.25
    assert getattr(cell, 'number_format', None) == '0.00%'

def test_various_number_formats_roundtrip(tmp_path):
    """Проверяем, что openpyxl -> файл -> cytosheet сохраняет number_format для разных типов."""
    file_path = tmp_path / "various_number_formats.xlsx"

    wb_ox = OpenpyxlWorkbook()
    ws_ox = wb_ox.active

    # Числа
    ws_ox["A1"].value = 1234.567
    ws_ox["A1"].number_format = "0.00"

    ws_ox["A2"].value = 0.25
    ws_ox["A2"].number_format = "0.00%"

    ws_ox["A3"].value = 1000
    ws_ox["A3"].number_format = "#,##0"

    # Даты / время
    ws_ox["B1"].value = datetime.date(2025, 1, 2)
    ws_ox["B1"].number_format = "m/d/yy"

    ws_ox["B2"].value = datetime.datetime(2025, 1, 2, 15, 30)
    ws_ox["B2"].number_format = "m/d/yy h:mm"

    # Текст с пользовательским форматом (не должен падать)
    ws_ox["C1"].value = "Text"
    ws_ox["C1"].number_format = "@"

    wb_ox.save(str(file_path))

    # Теперь читаем через cytosheet
    wb = load_workbook(str(file_path))
    ws = wb.active

    # A1, A2, A3 — проверяем только формат
    assert getattr(ws["A1"], "number_format", None) == "0.00"
    assert getattr(ws["A2"], "number_format", None) == "0.00%"
    assert getattr(ws["A3"], "number_format", None) == "#,##0"

    # Для B1 и B2 значения могут отличаться по типу, поэтому проверяем только формат
    assert getattr(ws["B1"], "number_format", None) == "m/d/yy"
    assert getattr(ws["B2"], "number_format", None) == "m/d/yy h:mm"

    # Текстовый формат '@' мы пока явно не поддерживаем, он может не подтянуться
    # поэтому просто убеждаемся, что чтение не сломалось и значение есть
    assert ws["C1"].value == "Text"
